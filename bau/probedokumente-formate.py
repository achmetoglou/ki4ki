#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Probedokumente in den ANDEREN Formaten, die die Aufnahme kennt.

Die Klassifizierung in Ablaufplan 1 kennt:

    pdf                        -> direkt zu Docling
    xlsx, xls                  -> eigener Excel-Weg
    doc, docx, odt, rtf        -> ERST ueber den Office-Dienst nach PDF
    ppt, pptx, odp             -> ebenso
    csv, txt, html, htm        -> Textwege
    alles andere               -> 'unsupported'

Diese Probe deckt vier davon ab. Was sie je einzeln beantworten soll:

  .docx   Das Dokument ist das ORIGINAL, die gewandelte PDF liegt daneben.
          Bekaeme die PDF einen eigenen Abdruck, spraenge jeder Beleg des
          Word-Dokuments ins Leere - lautlos. (bau/probedokument-word.js)
  .xlsx   Kommen Zahlen aus einer Tabelle ueberhaupt in der Antwort an?
          ⛔ OHNE Formeln: openpyxl schreibt Formeln ohne zwischen-
          gespeichertes Ergebnis, und ohne LibreOffice kann sie hier
          niemand berechnen. Die Zelle waere fuer die Aufnahme leer - die
          Probe wuerde dann etwas anderes pruefen, als sie behauptet.
  .txt    Der kuerzeste Weg, ohne Wandlung und ohne Seitenbegriff.
  .msg    ⭐ Der ABWEISUNGSWEG. Outlook-Nachrichten liegen in echten
          Kundenordnern haeufig herum, die Anlage kennt das Format nicht.
          Sie muss in aussortiert/ landen, MIT Begruendung - nicht still
          im Archiv. (Bekannte Schwachstelle: leerer Text landete bisher
          trotzdem im Archiv.)
          ⚠ Der Inhalt ist gewoehnliche MIME-Post, kein echtes
          Outlook-Format. Das ist hier richtig so: Die Klassifizierung
          entscheidet allein an der ENDUNG und liest die Datei nie. Eine
          echte .msg naehme denselben Weg.

Jede Zahl kommt genau einmal vor - im ganzen Probebestand.

  ~/.local/share/skills-venv/bin/python bau/probedokumente-formate.py <ziel>
"""
import io
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

SCHRIFT = "Arial"

TABELLE = [
    ("Kennwert", "Wert", "Einheit", "Prüfnorm"),
    ("Schmelztemperatur", "263", "°C", "DIN EN ISO 11357"),
    ("Dichte", "1,41", "g/cm³", "DIN EN ISO 1183"),
    ("Wasseraufnahme bei Sättigung", "2,8", "%", "DIN EN ISO 62"),
    ("Schwindung längs", "0,4", "%", "DIN EN ISO 294-4"),
]

TEXT = """Lieferantenliste Wareneingang

Die Sperrfrist für beanstandete Chargen beträgt 14 Tage.
Rückfragen gehen an die Qualitätssicherung.
Änderungen an dieser Liste sind zu dokumentieren.

Erfundenes Probedokument - keine echten Daten.
"""

POST = """From: qs@example.invalid
To: einkauf@example.invalid
Subject: Terminabsprache Wareneingang
Date: Tue, 22 Sep 2026 09:15:00 +0200
MIME-Version: 1.0
Content-Type: text/plain; charset=utf-8

Guten Tag,

für die Prüfung der nächsten Lieferung schlagen wir Donnerstag vor.

Erfundenes Probedokument - keine echten Daten.
"""


def tabelle_bauen(pfad):
    wb = Workbook()
    b = wb.active
    b.title = "Kennwerte"
    for zeile in TABELLE:
        b.append(list(zeile))
    for zelle in b[1]:
        zelle.font = Font(name=SCHRIFT, bold=True)
    for zeile in b.iter_rows(min_row=2):
        for zelle in zeile:
            zelle.font = Font(name=SCHRIFT)
    for spalte, breite in zip("ABCD", (30, 10, 10, 24)):
        b.column_dimensions[spalte].width = breite
    # Woher die Zahlen kommen, dort wo es jemand liest.
    b.cell(row=len(TABELLE) + 2, column=1,
           value="Erfundene Werte für einen Aufnahme-Test – keine echten Daten."
           ).font = Font(name=SCHRIFT, italic=True, size=9)
    os.makedirs(os.path.dirname(pfad), exist_ok=True)
    wb.save(pfad)
    return pfad


def text_bauen(pfad, inhalt):
    os.makedirs(os.path.dirname(pfad), exist_ok=True)
    with io.open(pfad, "w", encoding="utf-8") as fh:
        fh.write(inhalt)
    return pfad


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    ziel = sys.argv[1]
    gebaut = [
        tabelle_bauen(os.path.join(
            ziel, "Müller & Söhne", "Kennwerte", "Werkstoffkennwerte.xlsx")),
        text_bauen(os.path.join(
            ziel, "Schäfer Kunststofftechnik", "Qualität", "Wareneingang",
            "Lieferantenliste.txt"), TEXT),
        text_bauen(os.path.join(
            ziel, "Schäfer Kunststofftechnik", "Schriftverkehr",
            "Terminabsprache.msg"), POST),
    ]
    for p in gebaut:
        print("  %6d Byte  %s" % (os.path.getsize(p), os.path.relpath(p, ziel)))
    print("\nDas Word-Dokument kommt aus bau/probedokument-word.js "
          "(docx-js, eigener Lauf).")
    print("""
Fragen, die zu stellen sind (jede Zahl kommt genau einmal vor):
  - "Welche Schmelztemperatur ist angegeben?"        -> 263 °C   (Excel)
  - "Wie lange ist die Sperrfrist?"                  -> 14 Tage  (Text)
  - "Welche Zugfestigkeit hat Charge 11?"            -> 371 MPa  (Word)

⛔ Und die Probe, die schiefgehen DARF:
  Die .msg-Datei muss in aussortiert/ landen, mit Begruendung. Liegt sie
  im Archiv, ist der Abweisungsweg offen - dann kaeme bei KAP jede
  Outlook-Nachricht als leeres Dokument in den Bestand.""")
    return 0


if __name__ == "__main__":
    sys.exit(main())
